"""Local input adapters for editable DOCX or text-derived rewrite drafts."""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path
from typing import Any

from pypdf import PdfReader

from journalport.transform.hashing import file_hash

SUPPORTED = (".docx", ".md", ".txt", ".tex", ".pdf")
HEADINGS = re.compile(
    r"^(?:(?:[IVXLCDM]+|\d+(?:\.\d+)*)[.)]?\s+)?(abstract|introduction|related works?|background|methods?|methodology|materials and methods|results|experiments|experimental results|discussion|conclusions?|references|bibliography|acknowledg[e]?ments|author contributions|competing interests|data availability|code availability)\s*:?$",
    re.IGNORECASE,
)


def prepare_rewrite_input(source: Path, directory: Path) -> tuple[Path, dict[str, Any]]:
    """Convert locally, archive original and disclose conversion fidelity limitations."""
    source = source.resolve()
    if source.suffix.lower() not in SUPPORTED:
        raise ValueError(
            "Supported manuscript inputs: DOCX, Markdown, text, LaTeX and text-based PDF"
        )
    if source.stat().st_size > 100_000_000:
        raise ValueError("Input exceeds 100 MB")
    directory.mkdir(parents=True, exist_ok=False)
    directory.chmod(0o700)
    archived = directory / ("original" + source.suffix.lower())
    shutil.copyfile(source, archived)
    archived.chmod(0o600)
    report: dict[str, Any] = {
        "input_format": source.suffix.lower(),
        "source_sha256": file_hash(source),
        "original": str(archived),
        "warnings": [],
        "lossless_conversion": False,
    }
    normalized = directory / "normalized.docx"
    suffix = source.suffix.lower()
    if suffix == ".docx":
        shutil.copyfile(source, normalized)
        report["lossless_conversion"] = True
    elif suffix in (".md", ".tex"):
        try:
            import pypandoc  # type: ignore[import-untyped]
        except ImportError as exc:
            raise ValueError(
                "Markdown/LaTeX conversion needs pip install 'journalport[rewrite]'"
            ) from exc
        # Do not execute TeX or arbitrary user filters. Pandoc performs parsing only.
        original = source.read_text()
        if suffix == ".tex" and re.search(
            r"\\(?:input|include|write18|openout|read|write)\b", original
        ):
            raise ValueError(
                "LaTeX includes or file-I/O commands require a self-contained manuscript"
            )
        if suffix == ".md" and re.search(r"!\[[^]]*\]\((?:https?://|/|\.\.)", original):
            raise ValueError("Markdown images must be local files within the manuscript directory")
        pypandoc.convert_file(
            str(source),
            "docx",
            format="markdown" if suffix == ".md" else "latex",
            outputfile=str(normalized),
            extra_args=["--standalone", "--sandbox", "--resource-path=" + str(source.parent)],
        )
        report["warnings"] = [
            "Pandoc conversion requires formula, figure and bibliography inspection; original is archived."
        ]
    else:
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError(
                "Text/PDF conversion needs pip install 'journalport[rewrite]'"
            ) from exc
        if suffix == ".pdf":
            reader = PdfReader(source)
            if reader.is_encrypted:
                raise ValueError("Decrypt PDF before rewriting")
            pages = [page.extract_text(extraction_mode="layout") or "" for page in reader.pages]
            if any(len(page.strip()) < 30 for page in pages):
                raise ValueError(
                    "PDF has scanned/empty pages; provide an editable manuscript or checked OCR"
                )
            value = "\n\n".join(pages)
            report["warnings"] = [
                "PDF text extraction cannot guarantee reading order or recover editable formulas/tables/figures. Original PDF is archived; visual source comparison required."
            ]
        else:
            value = source.read_text(encoding="utf-8-sig")
        doc = Document()
        if not any(HEADINGS.fullmatch(line.strip()) for line in value.splitlines()):
            doc.add_heading("Source manuscript", level=1)
        for block in re.split(r"\n\s*\n", value):
            for segment in re.split(
                r"\n(?=(?:(?:[IVXLCDM]+|\d+)[.)]?\s+)?(?:Abstract|Introduction|Methods|Results|Discussion|Conclusion|References)\b)",
                block.strip(),
            ):
                lines = segment.strip().splitlines()
                if not lines:
                    continue
                if HEADINGS.fullmatch(lines[0].strip()):
                    doc.add_heading(lines[0].strip(), level=1)
                    if len(lines) > 1:
                        doc.add_paragraph(" ".join(lines[1:]))
                else:
                    doc.add_paragraph(" ".join(lines))
        doc.save(str(normalized))
    normalized.chmod(0o600)
    report["normalized_sha256"] = file_hash(normalized)
    (directory / "conversion_report.json").write_text(json.dumps(report, indent=2) + "\n")
    return normalized, report
